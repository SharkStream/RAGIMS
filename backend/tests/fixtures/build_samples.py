"""Sample file builders for the parsing pipeline tests (D5 step 1).

Two roles in one module:

1. ``*_bytes`` functions — deterministic in-memory sample builders used by
   the offline test suite (the DSH sandbox forbids subprocess writes, so
   tests never rely on files on disk).
2. ``main()`` — writes the same samples as real files into ``tests/fixtures/``
   (``sample.pdf/sample.docx/sample.txt``) for humans and external tools.
   Run manually from ``backend/``::

       env/Scripts/python.exe tests/fixtures/build_samples.py

Keep the builders deterministic: identical bytes on every call (fixed seed
word lists, no timestamps) so hashes and chunk counts are stable.
"""

from __future__ import annotations

import io
import sys
from pathlib import Path

import pymupdf
from docx import Document as DocxDocument

_WORDS = ["alpha", "beta", "gamma", "delta", "epsilon", "zeta", "eta", "theta"]


def lorem(tokens_hint: int = 60, offset: int = 0) -> str:
    """Deterministic filler text of roughly ``tokens_hint`` cl100k tokens."""
    out: list[str] = []
    i = offset
    while len(" ".join(out)) < tokens_hint * 4:  # ~4 chars per token for these words
        out.append(_WORDS[i % len(_WORDS)])
        i += 1
    return " ".join(out)


def build_pdf_bytes(*, header: str | None = None, pages: int = 3) -> bytes:
    """Multi-page PDF; one text line per insert (insert_text never wraps)."""
    doc = pymupdf.open()
    try:
        doc.set_metadata({"title": "RAGIMS Sample", "author": "ragims"})
        for i in range(pages):
            page = doc.new_page()
            y = 72.0
            if header:
                page.insert_text((72, y), header)
                y += 15
            page.insert_text((72, y), f"Sample section {i + 1} heading")
            y += 15
            for j in range(6):
                line = f"p{i + 1} l{j} {lorem(20, offset=i * 3 + j)}"[:85]
                page.insert_text((72, y), line)
                y += 15
        return doc.tobytes()
    finally:
        doc.close()


def build_encrypted_pdf_bytes(password: str = "pw") -> bytes:
    doc = pymupdf.open()
    try:
        doc.new_page().insert_text((72, 72), "top secret content")
        return doc.tobytes(
            encryption=pymupdf.PDF_ENCRYPT_AES_256, owner_pw=password, user_pw=password
        )
    finally:
        doc.close()


def build_docx_bytes() -> bytes:
    doc = DocxDocument()
    doc.add_heading("RAGIMS Sample Report", level=1)
    doc.add_paragraph(lorem(80))
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "key", "value"
    table.cell(1, 0).text, table.cell(1, 1).text = "alpha", "42"
    doc.add_paragraph(lorem(60, offset=3))
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def build_txt_bytes() -> bytes:
    lines = [f"line {i}: {lorem(15, offset=i)}" for i in range(10)]
    return ("\n".join(lines) + "\n").encode("utf-8")


SAMPLES = {
    "sample.pdf": lambda: build_pdf_bytes(header="Q3 2026 Confidential"),
    "sample.docx": build_docx_bytes,
    "sample.txt": build_txt_bytes,
}


def main() -> int:
    """Write the sample files next to this script (manual convenience)."""
    here = Path(__file__).resolve().parent
    for name, builder in SAMPLES.items():
        target = here / name
        try:
            target.write_bytes(builder())
        except PermissionError:
            print(
                f"[skip] {name}: sandbox forbids writing {target} — tests build "
                f"samples in memory, so this only affects having real files on disk",
                file=sys.stderr,
            )
            continue
        print(f"[ok] wrote {target} ({target.stat().st_size} bytes)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
