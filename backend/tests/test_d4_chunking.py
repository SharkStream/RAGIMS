"""D4 chunker tests (offline, tiktoken encoding is cached locally).

Run from ``backend/``::

    env/Scripts/python.exe -m pytest tests/test_d4_chunking.py -v -p no:cacheprovider
"""

from __future__ import annotations

import io

import pymupdf
import tiktoken
from app.chunkers.base import (
    DEFAULT_CHUNKER_CONFIG,
    BaseChunker,
    Chunk,
    ChunkerConfig,
)
from app.chunkers.fixed_token_chunker import FixedTokenChunker
from app.parsers.base import Block, ParsedDocument
from app.parsers.docx_parser import DocxParser
from app.parsers.pdf_parser import PdfParser
from app.parsers.txt_parser import TxtParser
from app.services.cleaning import clean
from app.utils.hashing import sha256_hex
from docx import Document as DocxDocument

ENC = tiktoken.get_encoding("cl100k_base")
chunker = FixedTokenChunker()


def ntok(text: str) -> int:
    return len(ENC.encode(text))


def doc(*blocks: Block) -> ParsedDocument:
    return ParsedDocument(text="\n\n".join(b.text for b in blocks), blocks=list(blocks))


def long_text(tokens: int) -> str:
    """Deterministic text with roughly ``tokens`` cl100k tokens."""
    words = ["alpha", "beta", "gamma", "delta", "epsilon", "zeta", "eta", "theta"]
    out: list[str] = []
    while ntok(" ".join(out)) < tokens:
        out.append(words[len(out) % len(words)])
    return " ".join(out)


# --------------------------------------------------------------------------- #
# config (step 3)
# --------------------------------------------------------------------------- #


def test_default_config_matches_mvp_spec() -> None:
    assert DEFAULT_CHUNKER_CONFIG.chunk_size == 512
    assert DEFAULT_CHUNKER_CONFIG.overlap == 50


def test_config_rejects_invalid_values() -> None:
    import pytest

    for size, overlap in [(0, 0), (512, 512), (512, 600), (512, -1)]:
        with pytest.raises(ValueError):
            ChunkerConfig(chunk_size=size, overlap=overlap)


def test_base_chunker_is_abstract() -> None:
    import pytest

    with pytest.raises(TypeError):
        BaseChunker()  # type: ignore[abstract]


# --------------------------------------------------------------------------- #
# core behaviour
# --------------------------------------------------------------------------- #


def test_short_block_yields_single_chunk() -> None:
    chunks = chunker.chunk(doc(Block(text="short text", page=1)))
    assert len(chunks) == 1
    assert chunks[0].chunk_index == 0
    assert chunks[0].text == "short text"
    assert chunks[0].page == 1
    assert chunks[0].token_count == ntok("short text")


def test_long_block_is_windowed_by_chunk_size() -> None:
    cfg = ChunkerConfig(chunk_size=50, overlap=10)
    body = long_text(400)
    chunks = chunker.chunk(doc(Block(text=body, page=3)), cfg)

    assert len(chunks) > 1
    assert [c.chunk_index for c in chunks] == list(range(len(chunks)))
    # every full chunk is exactly chunk_size; the tail is <= chunk_size
    for c in chunks[:-1]:
        assert c.token_count == cfg.chunk_size
    assert chunks[-1].token_count <= cfg.chunk_size
    # acceptance: every chunk carries more than `overlap` tokens
    for c in chunks:
        assert cfg.overlap < c.token_count <= cfg.chunk_size


def test_overlap_repeats_boundary_tokens() -> None:
    cfg = ChunkerConfig(chunk_size=30, overlap=10)
    chunks = chunker.chunk(doc(Block(text=long_text(200))), cfg)
    first = ENC.encode(chunks[0].text)
    second = ENC.encode(chunks[1].text)
    # the last `overlap` tokens of chunk 0 open chunk 1
    assert first[-cfg.overlap :] == second[: cfg.overlap]


def test_zero_overlap_splits_without_repeat() -> None:
    cfg = ChunkerConfig(chunk_size=20, overlap=0)
    chunks = chunker.chunk(doc(Block(text=long_text(100))), cfg)
    assert all(c.token_count == 20 for c in chunks[:-1])
    assert ENC.encode(chunks[0].text)[-1] != ENC.encode(chunks[1].text)[0]


def test_chunks_never_span_blocks_and_keep_pages() -> None:
    cfg = ChunkerConfig(chunk_size=40, overlap=10)
    blocks = [
        Block(text=long_text(120), page=1),
        Block(text=long_text(120), page=2, block_type="table"),
        Block(text="tail page three", page=3),
    ]
    chunks = chunker.chunk(doc(*blocks), cfg)

    pages = [c.page for c in chunks]
    assert pages == sorted(pages)  # block order preserved
    assert set(pages) == {1, 2, 3}
    assert chunks[-1].text == "tail page three"
    # table block chunks carry the block type through metadata
    assert any(c.metadata["block_type"] == "table" for c in chunks)


def test_chunk_hash_is_sha256_of_text() -> None:
    chunks = chunker.chunk(doc(Block(text=long_text(100))), ChunkerConfig(60, 10))
    for c in chunks:
        assert c.chunk_hash == sha256_hex(c.text.encode("utf-8"))
        assert len(c.chunk_hash) == 64


def test_tail_window_stays_within_bounds() -> None:
    # Sweep block lengths around the window boundary: every emitted chunk must
    # carry more than `overlap` tokens and at most `chunk_size`.
    cfg = ChunkerConfig(chunk_size=50, overlap=10)
    for n in (48, 50, 51, 55, 80, 81, 82, 90, 91, 130):
        chunks = chunker.chunk(doc(Block(text=long_text(n))), cfg)
        assert chunks, f"len={n}: no chunks"
        for c in chunks:
            assert (
                cfg.overlap < c.token_count <= cfg.chunk_size
            ), f"len={n}: token_count={c.token_count}"


def test_does_not_mutate_input() -> None:
    block = Block(text=long_text(100), page=1)
    parsed = doc(block)
    snapshot = block.text
    chunker.chunk(parsed, ChunkerConfig(30, 5))
    assert parsed.blocks[0] is block and block.text == snapshot


def test_empty_document_yields_no_chunks() -> None:
    assert chunker.chunk(ParsedDocument(text="", blocks=[])) == []


def test_page_is_none_for_txt_blocks() -> None:
    parsed = ParsedDocument(text="hello world", blocks=[Block(text="hello world")])
    chunks = chunker.chunk(parsed)
    assert chunks[0].page is None


def test_chunk_dataclass_defaults() -> None:
    c = Chunk(chunk_index=0, text="t", token_count=1)
    assert c.page is None and c.metadata == {} and c.chunk_hash == ""


# --------------------------------------------------------------------------- #
# acceptance: all three real formats -> cleaned -> non-empty chunks
# --------------------------------------------------------------------------- #


def build_pdf() -> bytes:
    d = pymupdf.open()
    try:
        for i in range(3):
            page = d.new_page()
            page.insert_text((72, 72), f"Section {i + 1} overview")
            # insert_text never wraps: emit short lines inside the page width
            for j in range(8):
                line = (f"p{i + 1} l{j} " + long_text(30))[:85]
                page.insert_text((72, 92 + j * 15), line)
        return d.tobytes()
    finally:
        d.close()


def build_docx() -> bytes:
    d = DocxDocument()
    d.add_heading("Chunk Source", level=1)
    d.add_paragraph(long_text(200))
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "k", "v"
    t.cell(1, 0).text, t.cell(1, 1).text = "a", "b"
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def test_all_three_formats_produce_non_empty_chunks() -> None:
    cfg = ChunkerConfig(chunk_size=60, overlap=10)
    sources = {
        "pdf": PdfParser().parse_bytes(build_pdf()),
        "docx": DocxParser().parse_bytes(build_docx()),
        "txt": TxtParser().parse_bytes(long_text(300).encode()),
    }
    for file_type, parsed in sources.items():
        cleaned = clean(parsed)
        chunks = chunker.chunk(cleaned, cfg)
        assert chunks, f"{file_type}: no chunks"
        assert all(c.text.strip() for c in chunks), f"{file_type}: empty chunk text"
        assert all(c.chunk_hash for c in chunks), f"{file_type}: missing hash"
        assert [c.chunk_index for c in chunks] == list(range(len(chunks)))
        for c in chunks:
            assert 0 < c.token_count <= cfg.chunk_size
        # long blocks honour the (overlap, chunk_size] window bound; short
        # blocks (e.g. a docx heading) are emitted whole by design
        if file_type != "docx":
            assert all(c.token_count > cfg.overlap for c in chunks)
        print(f"  {file_type}: {len(cleaned.blocks)} blocks -> {len(chunks)} chunks")
