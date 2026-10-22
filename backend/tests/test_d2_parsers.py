"""D2 parser tests (offline — all three formats built in memory).

* PDF bytes via PyMuPDF itself (new doc -> pages -> text -> tobytes).
* DOCX bytes via python-docx itself (headings, paragraphs, table).
* TXT bytes: utf-8, GBK (encoding detection), CRLF normalization, empty.

``parse_and_store`` (MinIO write + status writes) is DB-coupled and verified
by the live probe recorded in ``docs/tasks/d2-format-parsers.md``.

Run from ``backend/``::

    env/Scripts/python.exe -m pytest tests/test_d2_parsers.py -v -p no:cacheprovider
"""

from __future__ import annotations

import io

import pymupdf
from app.parsers.base import ParsedDocument
from app.parsers.docx_parser import DocxParser
from app.parsers.pdf_parser import PdfParser
from app.parsers.registry import get_parser, registered_types
from app.parsers.txt_parser import TxtParser
from app.services.parsing_service import parsed_object_key
from docx import Document as DocxDocument

# --------------------------------------------------------------------------- #
# in-memory file builders
# --------------------------------------------------------------------------- #


def build_pdf(pages: int = 2) -> bytes:
    doc = pymupdf.open()
    try:
        doc.set_metadata({"title": "C4 Report", "author": "ragims"})
        for i in range(pages):
            page = doc.new_page()
            page.insert_text((72, 72), f"page {i + 1} content")
        return doc.tobytes()
    finally:
        doc.close()


def build_docx() -> bytes:
    doc = DocxDocument()
    doc.add_heading("Report Title", level=1)
    doc.add_paragraph("First paragraph.")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "col a"
    table.cell(0, 1).text = "col b"
    table.cell(1, 0).text = "1"
    table.cell(1, 1).text = "2"
    doc.add_paragraph("After table.")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# --------------------------------------------------------------------------- #
# registry: real parsers registered
# --------------------------------------------------------------------------- #


def test_real_parsers_registered_for_all_three_types() -> None:
    import app.parsers.docx_parser  # noqa: F401 — self-registration
    import app.parsers.pdf_parser  # noqa: F401
    import app.parsers.txt_parser  # noqa: F401

    assert registered_types() == {"pdf", "docx", "txt"}
    assert isinstance(get_parser("pdf"), PdfParser)
    assert isinstance(get_parser("docx"), DocxParser)
    assert isinstance(get_parser("txt"), TxtParser)


# --------------------------------------------------------------------------- #
# PDF
# --------------------------------------------------------------------------- #


def test_pdf_parser_pages_blocks_and_metadata() -> None:
    parsed = PdfParser().parse_bytes(build_pdf(pages=3))
    assert isinstance(parsed, ParsedDocument)
    assert len(parsed.blocks) == 3
    assert [b.page for b in parsed.blocks] == [1, 2, 3]
    assert parsed.blocks[0].text.startswith("page 1")
    assert parsed.metadata.get("title") == "C4 Report"
    assert parsed.metadata.get("author") == "ragims"
    assert "page 2 content" in parsed.text


# --------------------------------------------------------------------------- #
# DOCX
# --------------------------------------------------------------------------- #


def test_docx_parser_headings_paragraphs_table_in_order() -> None:
    parsed = DocxParser().parse_bytes(build_docx())
    kinds = [b.block_type for b in parsed.blocks]
    assert kinds == ["heading", "paragraph", "table", "paragraph"]
    assert parsed.blocks[0].text == "Report Title"
    table_text = parsed.blocks[2].text
    assert table_text.startswith("| col a | col b |")
    assert "| --- | --- |" in table_text
    assert parsed.blocks[3].text == "After table."
    assert "First paragraph." in parsed.text


# --------------------------------------------------------------------------- #
# TXT
# --------------------------------------------------------------------------- #


def test_txt_parser_utf8() -> None:
    parsed = TxtParser().parse_bytes(b"line one\nline two")
    assert parsed.text == "line one\nline two"
    assert len(parsed.blocks) == 1
    assert parsed.metadata["encoding"].lower().startswith(("utf-8", "ascii"))


def test_txt_parser_detects_non_utf8_encoding() -> None:
    gbk_bytes = "中文内容测试".encode("gbk")
    parsed = TxtParser().parse_bytes(gbk_bytes)
    assert "中文内容测试" in parsed.text  # decoded, not mojibake


def test_txt_parser_normalizes_crlf() -> None:
    parsed = TxtParser().parse_bytes(b"a\r\nb\rc")
    assert parsed.text == "a\nb\nc"


def test_txt_parser_empty_file() -> None:
    parsed = TxtParser().parse_bytes(b"   \n  ")
    assert parsed.text == "" and parsed.blocks == []


def test_parse_template_reads_from_file_path() -> None:
    # Reuse this test file as the "on-disk" input for the template method.
    import pathlib

    parsed = TxtParser().parse(pathlib.Path(__file__))
    assert "build_docx" in parsed.text


# --------------------------------------------------------------------------- #
# object-key convention (D2 step 4)
# --------------------------------------------------------------------------- #


def test_parsed_object_key_convention() -> None:
    assert (
        parsed_object_key("12345678-1234-1234-1234-123456789012")
        == "parsed-text/12345678-1234-1234-1234-123456789012.json"
    )
